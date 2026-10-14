#!/usr/bin/env python3
"""
Two-way sync between the game's balance data and an editable spreadsheet.

  python scripts/balance_sheet.py export [file.xlsx]          code -> spreadsheet
  python scripts/balance_sheet.py import [file.xlsx]          dry run: show what would change
  python scripts/balance_sheet.py import [file.xlsx] --write  apply the changes to src/data/*.ts

Default file: attack-list-full.xlsx (project root). Needs: pip install openpyxl

What import can change
  Moves sheet          Type, Category, Power, Description
  Pet Learnsets sheet  which moves each pet learns and at what level (add/remove rows)
  Pets sheet           Type, Rarity, Base HP/Attack/Defense, Capture cost
  Weapons sheet        the moves a weapon grants
  Boss Specials sheet  the special moves a boss gets

Names are the lookup keys, so do NOT rename a move, pet, weapon or boss in the
sheet. New moves/pets must be added in the .ts files first, then re-exported.
Import validates everything first and writes nothing if any row is invalid.
"""
import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "src" / "data"
FILES = {k: DATA / f"{k}.ts" for k in ("moves", "pets", "items", "bosses")}
DEFAULT_XLSX = ROOT / "attack-list-full.xlsx"

TYPES = ["Normal", "Rock", "Nature", "Wood", "Water", "Light", "Shadow", "Fire", "Ice"]
CATEGORIES = ["Direct", "Ranged", "Chance"]
RARITIES = ["common", "uncommon", "rare"]
COLOURS = {"Normal": "9E9E9E", "Rock": "78909C", "Nature": "4CAF50", "Wood": "795548",
           "Water": "2196F3", "Light": "FFD54F", "Shadow": "4A148C", "Fire": "E53935", "Ice": "FFFFFF"}
LIGHT_TEXT = {"Light", "Ice"}  # types whose fill is bright enough to need dark text

# Top-level entries in these files are indented two spaces: `  'key': {` or `  key: {`.
BLOCK = re.compile(r"^  '?([\w-]+)'?: \{\n(.*?)^  \},?[ \t]*$", re.M | re.S)
ARRAY = lambda f: re.compile(rf"({f}: )\[([^\]]*)\]")
STRING_LIT = r"'((?:[^'\\]|\\.)*)'"


def read(key):
    return FILES[key].read_text(encoding="utf-8")


def str_items(raw):
    return re.findall(r"'([\w-]+)'", raw)


def unescape(s):
    return re.sub(r"\\(.)", r"\1", s)


def escape(s):
    return s.replace("\\", "\\\\").replace("'", "\\'")


def fmt_list(items):
    return "[" + ", ".join(f"'{i}'" if isinstance(i, str) else str(i) for i in items) + "]"


# ── Parsing ───────────────────────────────────────────────────────────────────

def parse_moves():
    out = {}
    for m in BLOCK.finditer(read("moves")):
        b = m.group(2)
        d = re.search(rf"description: {STRING_LIT}", b)
        out[m.group(1)] = dict(
            id=m.group(1),
            name=re.search(rf"name: {STRING_LIT}", b).group(1),
            type=re.search(r"type: PetType\.(\w+)", b).group(1),
            category=re.search(r"category: AttackCategory\.(\w+)", b).group(1),
            power=int(re.search(r"power: (\d+)", b).group(1)),
            description=unescape(d.group(1)) if d else "",
        )
    return out


def parse_pets():
    out = {}
    for m in BLOCK.finditer(read("pets")):
        b = m.group(2)
        g = lambda pat: re.search(pat, b)
        if not g(r"moveIds:"):
            continue
        out[m.group(1)] = dict(
            id=m.group(1),
            name=g(rf"name: {STRING_LIT}").group(1),
            type=g(r"type: PetType\.(\w+)").group(1),
            hp=int(g(r"baseHp: (\d+)").group(1)),
            atk=int(g(r"baseAttack: (\d+)").group(1)),
            defn=int(g(r"baseDefense: (\d+)").group(1)),
            moves=str_items(g(r"moveIds: \[([^\]]*)\]").group(1)),
            levels=[int(x) for x in re.findall(r"\d+", g(r"moveLevels: \[([^\]]*)\]").group(1))],
            cost=int(g(r"captureCost: (\d+)").group(1)),
            rarity=g(r"rarity: '(\w+)'").group(1),
        )
    return out


def parse_weapons():
    out = {}
    for m in BLOCK.finditer(read("items")):
        b = m.group(2)
        mv = re.search(r"moveIds: \[([^\]]*)\]", b)
        if mv:
            out[m.group(1)] = dict(id=m.group(1), name=re.search(rf"name: {STRING_LIT}", b).group(1),
                                   price=int(re.search(r"buyPrice: (\d+)", b).group(1)), moves=str_items(mv.group(1)))
    return out


def parse_bosses():
    out = {}
    for m in BLOCK.finditer(read("bosses")):
        b = m.group(2)
        sm = re.search(r"specialMoveIds: \[([^\]]*)\]", b)
        if sm:
            out[m.group(1)] = dict(id=m.group(1), name=re.search(rf"name: {STRING_LIT}", b).group(1),
                                   level=int(re.search(r"level: (\d+)", b).group(1)), moves=str_items(sm.group(1)))
    return out


# ── Export ────────────────────────────────────────────────────────────────────

def style_type_cell(c):
    if c.value in COLOURS:
        c.fill = PatternFill("solid", fgColor=COLOURS[c.value])
        c.font = Font(bold=True, color="000000" if c.value in LIGHT_TEXT else "FFFFFF")


def add_sheet(wb, title, head, rows, widths, type_cols=()):
    ws = wb.create_sheet(title)
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="333333")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for i in type_cols:
            style_type_cell(row[i])
    return ws


def dropdown(ws, col, values, n):
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{n + 1}")


def export(path):
    moves, pets, weapons, bosses = parse_moves(), parse_pets(), parse_weapons(), parse_bosses()
    names = {i: m["name"] for i, m in moves.items()}
    label = lambda mid: f"{names[mid]} ({moves[mid]['power']})" if mid in moves else f"{mid} (UNDEFINED)"
    by = lambda t: TYPES.index(t)
    weap_of, boss_of, pets_of = {}, {}, {}
    for w in weapons.values():
        for mid in w["moves"]:
            weap_of.setdefault(mid, []).append(f"{w['name']} (${w['price']})")
    for b in bosses.values():
        for mid in b["moves"]:
            boss_of.setdefault(mid, []).append(b["name"])
    for p in pets.values():
        for mid in p["moves"]:
            pets_of[mid] = pets_of.get(mid, 0) + 1

    wb = Workbook()
    wb.remove(wb.active)

    rows = [[m["name"], m["type"], "#" + COLOURS[m["type"]], m["category"], m["power"], m["description"],
             "; ".join(weap_of.get(m["id"], [])), "; ".join(boss_of.get(m["id"], [])), pets_of.get(m["id"], 0)]
            for m in sorted(moves.values(), key=lambda m: (by(m["type"]), m["power"]))]
    ws = add_sheet(wb, "Moves", ["Move", "Type", "Colour", "Category", "Power", "Description", "Weapon that grants it",
                                 "Used by bosses", "# pets that learn it"], rows, [18, 10, 10, 10, 8, 55, 28, 45, 12], (1,))
    dropdown(ws, "B", TYPES, len(rows))
    dropdown(ws, "D", CATEGORIES, len(rows))

    rows = []
    for p in pets.values():
        for mid, lv in zip(p["moves"], p["levels"]):
            m = moves.get(mid)
            rows.append([p["name"], p["type"], p["rarity"], names.get(mid, mid + " (UNDEFINED)"),
                         m["type"] if m else "?", m["category"] if m else "?", m["power"] if m else "?", lv])
    rows.sort(key=lambda r: (by(r[1]), r[0], r[7]))
    add_sheet(wb, "Pet Learnsets", ["Pet", "Pet type", "Rarity", "Move learned", "Move type", "Category", "Power",
                                    "Learned at level"], rows, [16, 10, 10, 18, 10, 10, 8, 15], (1, 4))

    rows = [[p["name"], p["type"], p["rarity"], p["hp"], p["atk"], p["defn"], p["cost"],
             ", ".join(f"{names.get(m, m)} Lv{l}" for m, l in zip(p["moves"], p["levels"]))] for p in pets.values()]
    rows.sort(key=lambda r: (by(r[1]), r[0]))
    ws = add_sheet(wb, "Pets", ["Pet", "Type", "Rarity", "Base HP", "Base Attack", "Base Defense", "Capture cost",
                                "Moves (read-only: edit on Pet Learnsets)"], rows, [16, 10, 10, 9, 11, 12, 13, 90], (1,))
    dropdown(ws, "B", TYPES, len(rows))
    dropdown(ws, "C", RARITIES, len(rows))

    rows = [[w["name"], w["price"]] + [label(m) for m in w["moves"]] for w in weapons.values()]
    add_sheet(wb, "Weapons", ["Weapon", "Price", "Move 1 (power)", "Move 2 (power)", "Move 3 (power)"], rows, [18, 8, 24, 24, 24])

    rows = [[b["name"], b["level"]] + [label(m) for m in b["moves"]] for b in bosses.values()]
    add_sheet(wb, "Boss Specials", ["Boss", "Level", "Special 1 (power)", "Special 2 (power)", "Special 3 (power)"],
              rows, [26, 7, 26, 26, 26])
    wb.save(path)
    print(f"Exported {len(moves)} moves, {len(pets)} pets, {len(weapons)} weapons, {len(bosses)} bosses -> {path}")


# ── Import ────────────────────────────────────────────────────────────────────

class Edits:
    """Collects regex edits per file and the human-readable change log."""

    def __init__(self):
        self.text = {k: read(k) for k in FILES}
        self.log, self.errors = [], []

    def err(self, msg):
        self.errors.append(msg)

    def set_in_block(self, file, key, field, old, new, repl):
        """Replace within one object block. `repl` is a (pattern, replacement) pair applied inside the block only."""
        if old == new:
            return
        for m in BLOCK.finditer(self.text[file]):
            if m.group(1) == key:
                body = re.sub(repl[0], repl[1], m.group(2), count=1)
                if body == m.group(2):
                    self.err(f"{key}: could not rewrite '{field}' in {file}.ts")
                    return
                t = self.text[file]
                self.text[file] = t[:m.start(2)] + body + t[m.end(2):]
                self.log.append(f"  {file}.ts  {key}.{field}: {old} -> {new}")
                return
        self.err(f"{key}: block not found in {file}.ts")


def cell_move(raw, name_to_id, where, edits):
    """'Rock Throw (15)' or 'Rock Throw' -> move id."""
    if raw is None or str(raw).strip() == "":
        return None
    name = re.sub(r"\s*\(\s*\d+\s*\)\s*$", "", str(raw)).strip()
    if name not in name_to_id:
        edits.err(f"{where}: unknown move '{name}' (names must match the Moves sheet exactly)")
        return None
    return name_to_id[name]


def rows_of(wb, title, edits):
    if title not in wb.sheetnames:
        edits.err(f"missing sheet '{title}'")
        return []
    return [r for r in wb[title].iter_rows(min_row=2, values_only=True) if any(c is not None for c in r)]


def to_int(v, where, edits, minimum=0):
    try:
        n = int(v)
        if n < minimum:
            raise ValueError
        return n
    except (TypeError, ValueError):
        edits.err(f"{where}: '{v}' must be a whole number >= {minimum}")
        return None


def import_sheet(path, write):
    wb = load_workbook(path, data_only=True)
    moves, pets, weapons, bosses = parse_moves(), parse_pets(), parse_weapons(), parse_bosses()
    e = Edits()
    name_to_id = {m["name"]: i for i, m in moves.items()}

    # Moves
    seen = set()
    for r in rows_of(wb, "Moves", e):
        name, typ, cat, power, desc = r[0], r[1], r[3], r[4], r[5]
        mid = name_to_id.get(name)
        if not mid:
            e.err(f"Moves: unknown move '{name}' (renaming/adding moves isn't supported)")
            continue
        seen.add(mid)
        m = moves[mid]
        if typ not in TYPES:
            e.err(f"Moves/{name}: type '{typ}' not one of {TYPES}")
            continue
        if cat not in CATEGORIES:
            e.err(f"Moves/{name}: category '{cat}' not one of {CATEGORIES}")
            continue
        p = to_int(power, f"Moves/{name} power", e)
        if p is None:
            continue
        desc = "" if desc is None else str(desc)
        e.set_in_block("moves", mid, "type", m["type"], typ, (r"type: PetType\.\w+", f"type: PetType.{typ}"))
        e.set_in_block("moves", mid, "category", m["category"], cat, (r"category: AttackCategory\.\w+", f"category: AttackCategory.{cat}"))
        e.set_in_block("moves", mid, "power", m["power"], p, (r"power: \d+", f"power: {p}"))
        if desc != m["description"]:
            esc = escape(desc)
            e.set_in_block("moves", mid, "description", repr(m["description"]), repr(desc),
                           (rf"description: {STRING_LIT}", lambda _m, esc=esc: f"description: '{esc}'"))
    for mid in set(moves) - seen:
        e.err(f"Moves: '{moves[mid]['name']}' row was deleted (delete the move in moves.ts instead)")

    # Pets (stats) -- keyed by name
    pet_by_name = {p["name"]: p for p in pets.values()}
    for r in rows_of(wb, "Pets", e):
        name = r[0]
        p = pet_by_name.get(name)
        if not p:
            e.err(f"Pets: unknown pet '{name}'")
            continue
        typ, rar = r[1], r[2]
        if typ not in TYPES:
            e.err(f"Pets/{name}: type '{typ}' invalid")
            continue
        if rar not in RARITIES:
            e.err(f"Pets/{name}: rarity '{rar}' not one of {RARITIES}")
            continue
        nums = [to_int(r[i], f"Pets/{name} {lbl}", e, 1 if i < 6 else 0)
                for i, lbl in ((3, "Base HP"), (4, "Base Attack"), (5, "Base Defense"), (6, "Capture cost"))]
        if None in nums:
            continue
        hp, atk, dfn, cost = nums
        pid = p["id"]
        e.set_in_block("pets", pid, "type", p["type"], typ, (r"type: PetType\.\w+", f"type: PetType.{typ}"))
        e.set_in_block("pets", pid, "rarity", p["rarity"], rar, (r"rarity: '\w+'", f"rarity: '{rar}'"))
        e.set_in_block("pets", pid, "baseHp", p["hp"], hp, (r"baseHp: \d+", f"baseHp: {hp}"))
        e.set_in_block("pets", pid, "baseAttack", p["atk"], atk, (r"baseAttack: \d+", f"baseAttack: {atk}"))
        e.set_in_block("pets", pid, "baseDefense", p["defn"], dfn, (r"baseDefense: \d+", f"baseDefense: {dfn}"))
        e.set_in_block("pets", pid, "captureCost", p["cost"], cost, (r"captureCost: \d+", f"captureCost: {cost}"))

    # Pet learnsets -- grouped per pet; rewritten only if the (move, level) set changed
    learn = {}
    for r in rows_of(wb, "Pet Learnsets", e):
        name = r[0]
        p = pet_by_name.get(name)
        if not p:
            e.err(f"Pet Learnsets: unknown pet '{name}'")
            continue
        mid = cell_move(r[3], name_to_id, f"Pet Learnsets/{name}", e)
        lv = to_int(r[7], f"Pet Learnsets/{name} level", e, 1)
        if mid and lv:
            if (mid, lv) in learn.setdefault(p["id"], []):
                e.err(f"Pet Learnsets/{name}: duplicate row {moves[mid]['name']} Lv{lv}")
            learn[p["id"]].append((mid, lv))
    for pid, p in pets.items():
        if pid not in learn:
            e.err(f"Pet Learnsets: '{p['name']}' has no rows (every pet needs at least one move)")
            continue
        new = learn[pid]
        if sorted(new) != sorted(zip(p["moves"], p["levels"])):
            ids, lvs = [m for m, _ in new], [l for _, l in new]
            old_s = ", ".join(f"{moves[m]['name']}@{l}" for m, l in zip(p["moves"], p["levels"]))
            new_s = ", ".join(f"{moves[m]['name']}@{l}" for m, l in new)
            e.set_in_block("pets", pid, "learnset", old_s, new_s, (r"moveIds: \[[^\]]*\],(\s*)moveLevels: \[[^\]]*\]",
                           lambda mm, a=fmt_list(ids), b=fmt_list(lvs): f"moveIds: {a},{mm.group(1)}moveLevels: {b}"))

    # Weapons / Bosses
    for title, entries, file, field, rx in (("Weapons", weapons, "items", "moveIds", r"moveIds: \[[^\]]*\]"),
                                            ("Boss Specials", bosses, "bosses", "specialMoveIds", r"specialMoveIds: \[[^\]]*\]")):
        by_name = {x["name"]: x for x in entries.values()}
        for r in rows_of(wb, title, e):
            x = by_name.get(r[0])
            if not x:
                e.err(f"{title}: unknown entry '{r[0]}'")
                continue
            ids = [cell_move(c, name_to_id, f"{title}/{r[0]}", e) for c in r[2:5]]
            ids = [i for i in ids if i]
            if not ids:
                e.err(f"{title}/{r[0]}: needs at least one move")
                continue
            if ids != x["moves"]:
                e.set_in_block(file, x["id"], field, ", ".join(x["moves"]), ", ".join(ids),
                               (rx, lambda _m, a=fmt_list(ids), f=field: f"{f}: {a}"))

    # Report
    if e.errors:
        print(f"{len(e.errors)} problem(s) found - nothing was changed:")
        print("\n".join("  ! " + x for x in e.errors))
        sys.exit(1)
    if not e.log:
        print("No differences between the spreadsheet and the code.")
        return
    print(f"{len(e.log)} change(s):")
    print("\n".join(e.log))
    if write:
        for k, f in FILES.items():
            f.write_text(e.text[k], encoding="utf-8")
        print("\nWritten. Run `npx tsc --noEmit` to confirm it still compiles.")
    else:
        print("\nDry run only. Re-run with --write to apply.")


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    flags = {a for a in sys.argv[1:] if a.startswith("--")}
    if not args or args[0] not in ("export", "import"):
        sys.exit(__doc__)
    path = Path(args[1]) if len(args) > 1 else DEFAULT_XLSX
    if args[0] == "export":
        export(path)
    else:
        if not path.exists():
            sys.exit(f"{path} not found - run `export` first")
        import_sheet(path, "--write" in flags)


if __name__ == "__main__":
    main()
